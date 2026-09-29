#!/usr/bin/env python3
"""One editable scene model -> Ecosphere-style HTML, PDF and native-text PPTX.

Use /tmp/ecosphere-deck-tools/bin/python. Quantitative atlas assets are generated
separately from the verified release. Never substitutes a synthetic atlas figure.
"""
from pathlib import Path
import argparse
import base64
import hashlib
import html
import json
import mimetypes
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from playwright.sync_api import sync_playwright

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'results/reports/ecosphere_grant_deck_20260929'
AS=OUT/'assets'
WHITE='FFFFFF'; OFF='E7E4DC'; MUT='AEC4C2'; ORANGE='E08A4A'; DARK='0A3C3F'; TEAL='78C6BB'
SERIF='Liberation Serif'; SANS='Liberation Sans'
SL=[]

def slide(kicker,title,source='',notes=''):
    s={'items':[],'notes':notes,'title':title};SL.append(s)
    text(s,88,58,1060,22,kicker.upper(),13,MUT,tracking=2.6)
    text(s,88,95,1104,112,title.upper(),42,WHITE,SERIF,line=1.10)
    rect(s,88,218,58,4,ORANGE)
    text(s,88,667,1058,29,source,11.5,MUT,line=1.2)
    text(s,1140,670,52,20,f'{len(SL):02}',13,MUT,align='right')
    return s

def text(s,x,y,w,h,value,size=22,color=OFF,font=SANS,bold=False,line=1.28,align='left',tracking=0,href=None):
    s['items'].append(dict(kind='text',x=x,y=y,w=w,h=h,value=value,size=size,color=color,font=font,bold=bold,line=line,align=align,tracking=tracking,href=href))
def rect(s,x,y,w,h,fill=DARK,stroke=None):
    s['items'].append(dict(kind='rect',x=x,y=y,w=w,h=h,fill=fill,stroke=stroke))
def line(s,x,y,w,h,color=ORANGE,width=2):
    s['items'].append(dict(kind='line',x=x,y=y,w=w,h=h,color=color,width=width))
def circle(s,x,y,d,fill=ORANGE):
    s['items'].append(dict(kind='circle',x=x,y=y,w=d,h=d,fill=fill))
def picture(s,x,y,w,h,path,fit='contain'):
    s['items'].append(dict(kind='image',x=x,y=y,w=w,h=h,path=str(path),fit=fit))
def tag(s,x,y,label,w=190):
    rect(s,x,y,w,28,ORANGE);text(s,x+10,y+6,w-20,20,label.upper(),12,DARK,bold=True,line=1)
def bullet(s,x,y,w,head,body):
    rect(s,x,y+8,6,6,ORANGE);text(s,x+22,y,w-22,30,head,22,WHITE,bold=True)
    text(s,x+22,y+35,w-22,72,body,19,MUT)
def box(s,x,y,w,h,head,body):
    rect(s,x,y,w,h,DARK,TEAL);text(s,x+20,y+20,w-40,34,head,25,WHITE,SERIF)
    text(s,x+20,y+65,w-40,h-70,body,19,MUT)


def make():
    # 01 — retain the original cover photograph, type hierarchy and coastal mood.
    s={'items':[],'title':'Ecosphere','notes':'Ecosphere Blue is the company; EmergentBiome is the platform; MethaNet is the first application. Greenhouse Gas Molecular Attestation Intelligence is the broader category. This deck presents a research funding proposal, not an awarded grant or an approved crediting method.'};SL.append(s)
    picture(s,0,0,1280,720,AS/'cover.jpg',fit='cover')
    # A pre-rendered transparent dark overlay preserves the original photo treatment.
    picture(s,0,0,1280,720,AS/'cover_overlay.png')
    text(s,88,196,1104,132,'ECOSPHERE',116,WHITE,SERIF,line=1)
    rect(s,88,347,58,4,ORANGE)
    text(s,88,377,700,93,'Make the biology behind\nwetland methane decision-ready.',33,OFF,line=1.24)
    text(s,88,508,680,65,'MethaNet, the first application of EmergentBiome\nGreenhouse Gas Molecular Attestation Intelligence',19,MUT)
    text(s,88,651,1104,25,'Alon Philosof, PhD  ·  Jayson Gutierrez, PhD  ·  Research funding proposal  ·  September 2026',14,MUT)

    # 02 — climate rationale without a generalized catastrophic-offset statistic.
    s=slide('The opportunity','Better restoration starts with\na clearer methane balance.',
        '[1] IPCC AR6 WGI, Ch. 7.  [2] Cotovicz et al., Nature Climate Change (2024).',
        'Approximately 80 times is the 20-year global warming potential of a methane pulse relative to an equal mass of CO2, not an atmospheric concentration or project offset. Cotovicz et al. report 10–33% reduction in potential aquatic methane emissions through surface-water oxidation in their study; this is neither a global wetland default nor Ecosphere performance. Both production and consumption, plus transport, determine emissions. Avoid the old 94–98% universal offset claim.')
    text(s,88,263,400,98,'~80×',82,ORANGE,SERIF,line=1)
    text(s,88,372,420,75,'Methane’s warming impact per unit mass\nrelative to CO₂ over 20 years.',22,WHITE)
    text(s,88,487,416,100,'A near-term climate reason to measure\nwhere methane is produced, consumed\nand released.',21,MUT)
    rect(s,574,264,618,323,DARK)
    text(s,610,290,550,30,'THE BIOLOGICAL OPPORTUNITY',15,TEAL,tracking=1.5)
    text(s,610,336,520,88,'10–33%',76,ORANGE,SERIF,line=1)
    text(s,610,439,534,115,'of potential aquatic methane emissions\nremoved by surface-water oxidation\nin a published mangrove study.',23,WHITE)
    text(s,88,614,1104,36,'Our question: can molecular evidence improve predictions of this balance?',24,WHITE)

    # 03 — reuse the landing-page artwork with explicit conceptual labeling.
    s=slide('The measurement layer','Connect microbial potential\nto measurements at the same place.',
        'Conceptual artwork from emergentbiome.earth. Pairing workflow proposed; the image is not field evidence.',
        'Gas measurements establish observed flux. DNA identifies potential; RNA provides evidence of transcription, not direct flux. Environmental covariates condition the interpretation. The site artwork is illustrative, not a photo of equipment installed at Cispata. Exact site, plot, depth, time and sample identity are essential. VM0033 already accounts for methane; the proposed additive contribution is molecular evidence and a test of its incremental value.')
    picture(s,570,244,622,410,AS/'field_cross_section.png')
    bullet(s,88,263,450,'Measure the outcome','Chamber methane flux, with time\nand spatial support recorded.')
    bullet(s,88,385,450,'Read the mechanism','DNA, selected RNA, pathway evidence\nand uncertainty in interpretation.')
    bullet(s,88,507,450,'Preserve the context','Salinity, sulfate, redox, inundation\nand carbon-stock measurements.')

    # 04 — company/platform/application hierarchy.
    s=slide('Ecosphere → EmergentBiome → MethaNet','One evidence platform.\nMethane is the first application.',
        'Architecture synthesized from the working MethaNet codebase and the public platform narrative.',
        'EmergentBiome integrates existing protein and genomic language models, functional annotations, source provenance and evidence review. It is not yet an independently trained general biome foundation model. MethaNet is the methane application. Molecular attestation means a traceable evidence packet and bounded interpretation, not statutory verification or carbon-credit approval. N2O and other ecosystems are future opportunities requiring their own validation.')
    text(s,88,264,387,111,'Greenhouse Gas\nMolecular Attestation\nIntelligence',29,WHITE,SERIF,line=1.23)
    text(s,88,414,385,106,'Turn molecular observations into\na traceable explanation, an explicit\nevidence gap and a next action.',23,MUT)
    tag(s,88,558,'Category ambition',220)
    box(s,552,254,640,112,'METHANET  /  FIRST APPLICATION','Methane screening and field-validation planning')
    line(s,872,366,0,27)
    box(s,552,393,640,147,'EMERGENTBIOME  /  SHARED PLATFORM','Protein + genomic representations · functional evidence\nSource provenance · review objects · reproducible analysis')
    line(s,872,540,0,24)
    text(s,552,579,640,53,'FUTURE: N₂O · other wetland systems · discovery\nEach expansion requires its own scientific validation.',18,MUT,align='center')

    # 05 — plotted only from the newly validated geometry.
    s=slide('What is built','A working molecular atlas,\nwith its evidence limits exposed.',
        'Molecular payload: 10 Aug 2026 ledger. Geometry: 29 Sep 2026 reconciliation. Counts are records, not independent samples.',
        '7,965 registered MAG/proteome records across four lanes; 7,710 have ESM2, gLM2 and functional views; 255 explicit gaps remain. 5,209 tri-view records use the shared functional pipeline and 2,501 use source annotations. None is fully mechanism-comparable across routes under the current gate. Pilot re-embedding fixes the identified ESM2 layer mismatch; it does not validate transfer, function, flux, or the graph as a predictor. The plot is a navigation figure only.')
    for y,n,l in [(265,'7,965','registered records'),(379,'7,710','three-view records'),(493,'255','explicit gaps retained')]:
        text(s,928,y,264,66,n,58,ORANGE,SERIF,line=1);text(s,928,y+71,264,28,l,19,MUT)
    mapfile=AS/'atlas_reconciled.png'
    if mapfile.exists():picture(s,88,251,778,347,mapfile)
    else:box(s,88,251,778,347,'ATLAS REBUILD IN PROGRESS','Final deck rendering requires the verified 29 September geometry.')
    text(s,88,620,1104,32,'Next gate: exact field pairing and independent tests of added predictive value.',22,WHITE)

    # 06 — actual source-bound product cases; no old nearest-neighbor claims.
    s=slide('MethaNet in use','Three evidence reviews.\nThree concrete next decisions.',
        'Public case projection, 26 Sep 2026. [3] MUCC v1. [4] Futian v3. Interpretation and follow-up: EmergentBiome.',
        'Cases are already present in the scoped landing-page demonstration. OWC_1706: CuMMO-family evidence requires methane-versus-ammonia substrate discrimination. MF1_201704_bin_10: five depth candidates are retained; exact library-to-MAG mapping remains unresolved. OWC_0000: MCR and RNA evidence support follow-up but do not establish methane direction or rate. Proposed next actions are not completed validations. Source licenses and transformations are in the figure/source ledger.')
    for x,num,pid,head,evidence,action in [
        (88,'01','OWC_1706','Resolve the substrate','CuMMO-family signal.\nMethane versus ammonia\ninterpretation remains open.','Targeted sequence review\nand substrate assay.'),
        (462,'02','MF1_201704_bin_10','Recover the field identity','One genome, five possible\nsediment depths. The\nambiguity stays visible.','Recover the original\nlibrary-to-depth crosswalk.'),
        (836,'03','OWC_0000','Test the process','MCR-family and RNA\nevidence. Direction and\nnet methane rate unproven.','Pair process assays\nwith chamber measurements.')]:
        rect(s,x,258,356,363,DARK);text(s,x+24,277,305,35,num,29,ORANGE,SERIF)
        text(s,x+24,323,309,25,pid,15,TEAL);text(s,x+24,366,309,60,head,28,WHITE,SERIF)
        text(s,x+24,432,309,83,evidence,19,MUT);line(s,x+24,532,306,0,TEAL,1)
        text(s,x+24,550,309,56,action,19,WHITE)

    # 07 — real experimental unit diagram.
    s=slide('The proposed field study','144 paired sample-events.\nA design that keeps context intact.',
        'Proposed design from the attached research application; site agreement and permits remain pending.',
        'Three restoration stages × four salinity positions × three plots = 36 plots. Each has two microsites, revisited in two seasons: 144 sample-events. Plots are the replicate unit, not 144 independent replicates. Salinity and restoration stage may be confounded in one hydrological unit, so causal restoration-effect claims require a suitable design beyond these observational contrasts. Cispata is a target at outreach stage; Tumaco or an equivalent site is a conditional fallback. Permits, access and benefit-sharing must precede fieldwork.')
    # Three rows, four columns, three plot markers per cell.
    for c in range(4):text(s,280+c*136,265,130,24,f'SALINITY {c+1}',13,MUT,align='center')
    for r,label in enumerate(['RECENTLY\nRECONNECTED','TRANSITIONAL','MATURE /\nREFERENCE']):
        y=307+r*90;text(s,88,y+13,176,55,label,16,WHITE)
        for c in range(4):
            x=284+c*136;rect(s,x,y,118,66,DARK)
            for k in range(3):circle(s,x+19+k*29,y+26,13,ORANGE)
    text(s,280,589,550,44,'Each dot = one plot. 36 plots total.',18,MUT)
    line(s,853,267,0,341,TEAL,1)
    text(s,899,269,277,67,'36 plots',44,WHITE,SERIF)
    text(s,899,349,277,49,'× 2 microsites',28,MUT,SERIF)
    text(s,899,413,277,49,'× 2 seasons',28,MUT,SERIF)
    text(s,899,491,277,85,'= 144',64,ORANGE,SERIF)
    text(s,899,578,277,51,'DNA + methane flux\n+ matched environment',19,WHITE)

    # 08 — locked tests, no invented performance bars.
    s=slide('The decisive experiment','Does molecular evidence add value\nbeyond the strongest baseline?',
        'Proposed evaluation plan. Freshwater rationale: [7] Bechtold et al., Nature Communications (2025). Saline transfer remains a test.',
        'Train and select hyperparameters within study/plot grouped folds. Use a pre-Cispata frozen model to test first-site transfer before incorporating any local labels. Use wet-season data to fit the local model; lock it and timestamp dry-season predictions before revealing the dry-season flux labels. Same-plot dry-season evaluation measures seasonal generalization at one site, not geographic generalization. Report MAE/RMSE, calibration/interval coverage, source-stratified errors and uncertainty clustered by plot. Handle signed flux and limits of detection explicitly; do not log-transform nonpositive observations blindly. Additional external-site validation remains necessary.')
    box(s,88,263,345,152,'01  ENVIRONMENT','Salinity · sulfate · redox\nHydrology · temperature\nRestoration context')
    box(s,468,263,345,152,'02  MARKER BASELINE','Environment + abundance\nof curated methane-cycle\nmarkers; RNA where available')
    box(s,848,263,344,152,'03  METHANET','The same baseline inputs\n+ molecular representations\nand curated pathway evidence')
    tag(s,88,455,'Freeze before testing',238)
    for x,a,b in [(88,'External data','Train with study-aware splits'),(469,'Wet-season field test','Test before local retraining'),(848,'Dry-season blind test','Lock predictions; reveal labels')]:
        text(s,x,510,344,33,a,26,WHITE,SERIF);text(s,x,558,344,56,b,20,MUT)
    line(s,376,492,726,0,ORANGE,2)

    # 09 — evidence ladder and buyer action.
    s=slide('From science to customer value','Help developers choose\nwhat to measure—and why.',
        'Commercial workflow is proposed. Current capability: molecular review and measurement planning.',
        'The beachhead is a scoped assessment with a restoration developer or research partner: ingest existing molecular and site data, deliver evidence packets and sampling priorities, then run a paired validation pilot. The intended economic value is better allocation of monitoring effort and clearer evidence for reviewers. Combine the methane measurements with measured carbon stocks and published scenario factors, reporting both 20- and 100-year horizons where climate-equivalence scenarios are presented. Two seasonal campaigns do not establish long-term carbon accumulation or permanence. Willingness to pay, delivery cost, repeat purchases and incremental predictive value need to be measured. Do not promise price uplift, cheaper tower-equivalent coverage, A–E ratings, or registry acceptance.')
    box(s,88,271,328,259,'REVIEW TODAY','What molecular evidence exists?\n\nWhere does interpretation stop?\n\nWhat measurement resolves it?')
    box(s,461,271,328,259,'VALIDATE NEXT','Does biology improve the\nmeasured methane estimate?\n\nWhere does the model fail?\n\nHow much uncertainty remains?')
    box(s,834,271,358,259,'DECIDE WITH EVIDENCE','Read methane alongside\ncarbon-stock evidence.\n\nCompare restoration scenarios\nand monitoring priorities\nwith explicit uncertainty.')
    text(s,88,582,1104,52,'First commercial milestone: a scoped, paid assessment with an agreed validation endpoint.',23,WHITE)

    # 10 — defensibility grounded in open science.
    s=slide('A platform that can compound','Open the funded evidence.\nBuild value in dependable delivery.',
        'Proposed grant outputs: SRA + Zenodo / CC BY 4.0; code and weights / Apache 2.0, subject to source and access terms.',
        'Do not describe the grant-funded paired dataset as proprietary; the application commits to open data, code, model and weights. Defensibility is an execution hypothesis: reliable sample provenance, harmonization, explicit abstention, validated models, workflow integration and repeat partner delivery. Separately consented future datasets and services may support commercial products. N2O, other ecosystems and a generative discovery layer are future avenues rather than current products or licensed revenue.')
    for x,n,h,b in [(88,'1','PAIRED OBSERVATIONS','Same place · depth · time\nMolecules + environment + flux'),(470,'2','REPRODUCIBLE EVIDENCE','Versioned inputs and models\nVisible uncertainty and failure'),(851,'3','TRUSTED WORKFLOW','Review packets and APIs\nPartner monitoring decisions')]:
        circle(s,x,273,51,ORANGE);text(s,x,282,51,35,n,27,DARK,SERIF,align='center')
        text(s,x,354,338,75,h,27,WHITE,SERIF);text(s,x,433,338,77,b,21,MUT)
    line(s,139,298,331,0,ORANGE,2);line(s,521,298,330,0,ORANGE,2)
    rect(s,88,558,1104,74,DARK)
    text(s,111,578,1058,44,'The reusable asset is the evidence system and its validation history.',25,WHITE,SERIF)

    # 11 — eight quarter proposed plan, gate driven.
    s=slide('Execution plan · proposed 2027–2028','Four phases.\nEach ends with a decision gate.',
        'Eight-quarter plan derived from the attached proposal; start date depends on funding and field access.',
        'Q1 permits/access and evidence inventory; Q2 preregister protocol and freeze pre-field model, site decision; Q3 wet-season campaign; Q4 pre-local test, local calibration and preprint; Q5 dry campaign with frozen local model; Q6 unblind tests and evaluate uncertainty; Q7 publish dataset/framework; Q8 peer review submission and final report. Access gate at Q2: activate an equivalent fallback site if needed. If molecular features add no value, publish the negative result and paired benchmark. Regulatory engagement is a future evidence review, not guaranteed methodology approval.')
    rows=[('Q1–Q2','READY THE STUDY','Permits, site agreement, source-aware baseline, frozen pre-field model.','Gate: access + preregistered protocol'),
          ('Q3–Q4','MEASURE & CALIBRATE','Wet-season field campaign; test transfer before local retraining.','Gate: paired data QC + baseline comparison'),
          ('Q5–Q6','TEST THE HELD-OUT SEASON','Dry-season campaign; timestamp predictions, then reveal flux labels.','Gate: added value + uncertainty calibration'),
          ('Q7–Q8','RELEASE & TRANSLATE','Open dataset, model and evaluation; restoration decision framework.','Gate: reproducible release + peer-review submission')]
    for i,(q,h,b,g) in enumerate(rows):
        y=257+i*97; text(s,88,y,146,49,q,34,ORANGE,SERIF)
        line(s,248,y+5,0,73,TEAL,1);text(s,271,y,331,32,h,23,WHITE,SERIF)
        text(s,271,y+38,825,30,b,19,MUT);text(s,660,y,522,30,g,17,TEAL)

    # 12 — qualifications from CV, avoid implied current partnerships.
    s=slide('The team','Biological interpretation\nand computational execution.',
        'Qualifications summarized from the investigators’ CVs in the supplied research application.',
        'Alon Philosof: microbial ecology, methanogen and viral genomics; Caltech postdoctoral research; expedition and single-cell characterization experience. Jayson Gutierrez: PhD at VIB/Ghent, postdoctoral research at Ghent and VLIZ; meta-omics, ecological modeling and the current EmergentBiome implementation. Listed institutions are prior affiliations, not endorsements. Site partners are at outreach stage; do not imply a signed INVEMAR/CI/CVS collaboration. Funded project includes local field staff and community engagement.')
    text(s,88,277,510,48,'ALON PHILOSOF, PhD',35,WHITE,SERIF)
    text(s,88,340,510,27,'PI · SCIENCE & FIELD PROGRAM',15,ORANGE,tracking=1.2)
    text(s,88,391,499,139,'Microbial ecology and methane-cycle genomics.\nCaltech postdoctoral research.\nMarine expedition and single-cell\ncharacterization experience.',23,OFF,line=1.48)
    line(s,640,275,0,287,TEAL,1)
    text(s,689,277,503,48,'JAYSON GUTIERREZ, PhD',35,WHITE,SERIF)
    text(s,689,340,503,27,'CO-PI · PLATFORM & MODELING',15,ORANGE,tracking=1.2)
    text(s,689,391,503,139,'Computational biology and meta-omics.\nVIB / Ghent PhD; Ghent and VLIZ\npostdoctoral research. Builder of the\nEmergentBiome atlas and inference system.',23,OFF,line=1.48)
    text(s,88,592,1104,52,'Local field teams and community agreements are part of the proposed field program.',21,MUT)

    # 13 — single coherent grant request.
    s=slide('The research ask','$450K · 24 months\nAn open test of molecular value.',
        'Proposed budget, not awarded funding. $215K personnel + $235K direct costs = $450K. No overhead/tax budgeted.',
        'The grant budget is separate from the older $650K pre-seed raise. Personnel $215K; two campaigns $60K; equipment $35K; laboratory $25K; DNA/RNA $45K; compute $40K; partner/community support $20K; release/publication $10K. Budget values are applicant estimates, not vendor quotes. Scope and feasibility should be confirmed at Q1. The official 2026 Google call closed September 25; this is a presentation/supporting deck, not a compliant substitute for its written application or evidence of submission/award.')
    items=[('People',215,ORANGE),('Field + equipment + lab',120,TEAL),('Sequencing',45,'BBCEC5'),('Compute',40,'699C99'),('Partners + open release',30,'D8D6C6')]
    y=268
    for label,n,col in items:
        text(s,88,y+1,345,28,label,19,OFF);rect(s,441,y,446*n/215,28,col);text(s,915,y,139,32,f'${n}K',24,WHITE,SERIF)
        y+=49
    text(s,88,553,1104,36,'Three deliverables worth funding',27,WHITE,SERIF)
    text(s,88,604,353,36,'01  Open paired benchmark',20,OFF)
    text(s,455,604,362,36,'02  Blind model evaluation',20,OFF)
    text(s,851,604,341,36,'03  Restoration decision guide',20,OFF)

    # Appendix 14 — maturity register, clear zero current claims.
    s=slide('Appendix A · evidence register','What is demonstrated.\nWhat remains to be established.',
        'Sources: release ledger, public case export, reconciliation record and the MRV risk-scoring roadmap.',
        'The 7,710 figure denotes view availability and schema normalization, not mechanism equivalence or sample validation. The current public atlas has zero final risk tiers and no exact sample-to-flux join supporting a calibrated sample score. These statuses are scientific gates, not a judgment that the underlying biology is absent. Retain missing and failed statuses in exports.')
    cols=[88,557,864];text(s,cols[0],265,449,25,'CAPABILITY',14,TEAL,tracking=1.5);text(s,cols[1],265,277,25,'EVIDENCE NOW',14,TEAL,tracking=1.5);text(s,cols[2],265,328,25,'NEXT GATE',14,TEAL,tracking=1.5)
    rows=[('Atlas and source provenance','7,965 registered records','Close 255 explicit gaps'),('Three molecular views','7,710 complete records','Cross-route comparability'),('Evidence packets','Three public review cases','Independent mechanism review'),('Exploratory ESM-2 geometry','Pilot pooling reconciled','Phylogeny/source-aware tests'),('Sample methane prediction','Not validated','Exact paired field evaluation'),('Risk tiers / crediting use','Not available','Calibration + external review')]
    for i,row in enumerate(rows):
        y=309+i*52;line(s,88,y-9,1104,0,'437672',.7)
        for c,t in enumerate(row):text(s,cols[c],y,[444,278,328][c],42,t,18,OFF if c==0 else MUT)

    # Appendix 15 — derived metrics inserted only after verification.
    reconciliation=OUT/'reconciliation_slide.json'
    r=json.loads(reconciliation.read_text()) if reconciliation.exists() else {}
    s=slide('Appendix B · configuration reconciliation','A shared representation\nrequires a shared configuration.',
        '662-proteome retained-FASTA recomputation; per-record hashes and reproduction checks. Scientific transfer remains unvalidated.',
        'The original pilot pooled layers 20–33; the later atlas lanes used layer 33. The recomputation generates both for each retained FASTA. Old-pooling vectors must reproduce the archived pilot within the recorded cosine tolerance before final-layer vectors are promoted. June lane final-layer attribution comes from archived run statistics, code history and numerical controls; their NPZ files do not all record model revisions. Remaining historical provenance limitations are explicit in the contract. No Procrustes alignment or ad hoc coordinate shift is substituted for recomputation.')
    text(s,88,268,301,38,'RECOMPUTE, THEN VERIFY',21,WHITE,SERIF)
    text(s,88,326,301,134,'Pilot: mean layers 20–33\nLater lanes: final layer 33\n\nRecomputed from the same\n662 retained proteomes.',19,MUT,line=1.4)
    controls=AS/'configuration_controls.png'
    if controls.exists():picture(s,426,247,766,240,controls)
    else:box(s,426,267,766,185,'CONTROL RETRIEVAL REBUILDING','Final plot requires the complete verified vector set.')
    text(s,88,496,1104,46,r.get('headline','Verification pending — not a release-ready slide.'),30,ORANGE,SERIF)
    text(s,88,565,1104,69,r.get('detail','The final deck must be regenerated after the GPU and verification jobs complete.'),19,MUT)

    # Appendix 16 — field rigor / fallbacks.
    s=slide('Appendix C · execution and scientific risks','A useful outcome\neven if the hypothesis fails.',
        'Proposed decision rules. Site-specific observational evidence does not establish universal transfer or long-term permanence.',
        'If molecular features do not beat strong environmental/marker baselines, release the paired data and negative result. Preserve adequate statistical power by adjusting sampling after preregistered design simulation; 144 events alone does not guarantee power. Restoration stage and salinity may covary; use blocked analysis and acknowledge confounding. Short campaigns do not establish long-term carbon accumulation or permanence: integrate historical monitoring, carbon stocks, sediment accretion and scenario uncertainty rather than claiming measured permanence. Proposed regulatory engagement should be evidence-led.')
    for y,h,b in [(266,'Access or permits delayed','Q2 site gate; activate a comparable fallback. Fieldwork starts after permissions.'),
                  (363,'No incremental molecular skill','Publish the benchmark, baseline comparison and negative result openly.'),
                  (460,'Poor transfer across sources or seasons','Stratify errors; narrow the supported domain and abstain outside it.'),
                  (557,'Carbon storage inferred too strongly','Use stock and context measurements; long-term permanence needs longer records.')]:
        text(s,88,y,405,74,h,26,WHITE,SERIF);text(s,552,y,640,66,b,21,MUT)

    # Appendix 17 — exact budget with quarter totals.
    s=slide('Appendix D · budget and delivery','One budget. Eight quarters.\nVisible assumptions.',
        'Attached proposal, Section 10. All values USD; estimates require operational confirmation. Total = $450,000.',
        'Quarter totals: 41875,76875,91875,36875,91875,31875,41875,36875. Personnel $215000 across 24 months. Direct costs $235000. No contingency, overhead or tax is explicitly budgeted in the source application; confirm institutional and tax treatment and vendor/field costs before committing expenditure. Do not conflate this grant budget with runway, investment valuation, or commercial gross margin.')
    costs=[('Personnel',215000),('Two field campaigns',60000),('Flux / porewater equipment',35000),('Laboratory analyses',25000),('DNA / selected RNA',45000),('Compute / storage',40000),('Partner / community support',20000),('Data release / publication',10000)]
    for i,(label,v) in enumerate(costs):
        y=263+i*42;text(s,88,y,365,30,label,19,OFF);text(s,453,y,150,30,f'${v:,}',21,WHITE,SERIF,align='right')
    line(s,653,262,0,339,TEAL,1)
    quarters=[41875,76875,91875,36875,91875,31875,41875,36875]
    for i,v in enumerate(quarters):
        x=704+i*60;h=230*v/max(quarters);rect(s,x,524-h,38,h,ORANGE)
        text(s,x-9,540,56,25,f'Q{i+1}',15,MUT,align='center')
        text(s,x-10,581,58,28,f'{v/1000:.1f}',14,WHITE,align='center')
    text(s,704,621,475,22,'Quarterly spend, USD thousands; labels rounded.',13,MUT)

    # Appendix 18 — readable selected bibliography; full ledger separate.
    s=slide('Appendix E · sources and use','Evidence behind the narrative.',
        'Verified 29 Sep 2026. Full URLs, claim corrections and figure provenance accompany the deck.',
        'Official Google call: up to $450000 for tidal wetland restoration research; closed September 25 2026. The requested deck is supporting presentation material. The call asks for maximum four overview pages plus one page per PI CV, six total with co-PI. The supplied PDF has seven pages; submission status is unknown. No application was submitted in this task. Original PDFs use Liberation Serif / Sans; these fonts and the original teal/copper design are retained. Field image is illustrative. Every quantitative figure is either a cited published result, a repo release count, or a labeled proposal.')
    refs=[('[1]','IPCC AR6 WGI, Chapter 7 (2021; updated tables).','Methane warming metric; time horizon explicitly stated.'),
          ('[2]','Cotovicz et al. Nature Climate Change 14, 275–281 (2024).','doi:10.1038/s41558-024-01927-1 · methane oxidation in mangroves'),
          ('[3]','Oliverio et al. MUCC, version 1.0.0.','doi:10.5281/zenodo.8194033 · CC BY 4.0'),
          ('[4]','Qi & Li. Futian seven-year genome catalogue, version 3.','doi:10.6084/m9.figshare.30883646.v3 · CC BY 4.0'),
          ('[5]','Verra VM0033 v2.1; v3.0 consultation, 3 Sep–5 Oct 2026.','Methane already addressed; a draft revision is not molecular approval.'),
          ('[6]','2026 Google Carbon Removal & Superpollutant R&D Awards.','Tidal-wetland priority · $450K ceiling · deadline 25 Sep 2026 (closed).'),
          ('[7]','Bechtold et al. Nature Communications 16, 944 (2025).','doi:10.1038/s41467-025-56133-0 · microbial interactions in freshwater wetlands')]
    urls=['https://www.ipcc.ch/report/ar6/wg1/chapter/chapter-7/',
          'https://doi.org/10.1038/s41558-024-01927-1',
          'https://doi.org/10.5281/zenodo.8194033',
          'https://doi.org/10.6084/m9.figshare.30883646.v3',
          'https://verra.org/methodologies/vm0033-methodology-for-tidal-wetland-and-seagrass-restoration-v2-1/',
          'https://www.research.google/programs-and-events/2026-google-carbon-removal-and-superpollutant-elimination-rd-awards/',
          'https://doi.org/10.1038/s41467-025-56133-0']
    for i,(n,a,b) in enumerate(refs):
        y=253+i*56;text(s,88,y,51,32,n,19,ORANGE);text(s,146,y,1046,28,a,20,WHITE,href=urls[i]);text(s,146,y+29,1046,25,b,16,MUT)

    # A direct product figure connects the conceptual cards to a working interface.
    graph=AS/'evidence_graph_live.png'
    if graph.exists():
        s=slide('Working product · source-linked evidence','Every conclusion stays attached\nto a record and a review decision.',
            'Live evidence explorer, captured 29 Sep 2026. MUCC v1 source facts. Links are authored review paths, not ecological interactions.',
            'This is the live OWC_1706 evidence-review interface. The four linked categories are recorded molecular evidence, interpretation on hold, unresolved sample context and a proposed next observation. The graph organizes a review; its positions and links are not learned ecological interactions, abundance estimates or methane rates. Source facts can be opened from the product. The user-facing value is a checkable decision trail.')
        picture(s,88,250,762,394,graph)
        text(s,897,272,295,33,'RECORDED',16,TEAL,tracking=1.5)
        text(s,897,309,295,60,'Four source genes\nwith traceable identities.',21,WHITE)
        text(s,897,405,295,33,'ON HOLD',16,ORANGE,tracking=1.5)
        text(s,897,442,295,61,'Substrate identity\nneeds further evidence.',21,WHITE)
        text(s,897,541,295,32,'NEXT ACTION',16,TEAL,tracking=1.5)
        text(s,897,578,295,60,'Review sequence context\nand test substrate use.',21,WHITE)
        SL.insert(6,SL.pop())
        for i,sl in enumerate(SL):
            for atom in sl['items']:
                if atom['kind']=='text' and atom['x']==1140 and atom['y']==670:atom['value']=f'{i+1:02}'


def uri(p):
    p=Path(p);return 'data:'+str(mimetypes.guess_type(p)[0] or 'application/octet-stream')+';base64,'+base64.b64encode(p.read_bytes()).decode()

def html_deck():
    fonts=[]
    for name,folder,stem in [(SERIF,'liberation-serif','LiberationSerif'),(SANS,'liberation-sans','LiberationSans')]:
        for style,weight in [('Regular',400),('Bold',700)]:
            file=Path('/usr/share/fonts')/folder/f'{stem}-{style}.ttf'
            if not file.exists():raise FileNotFoundError(f'Reference deck font missing: {file}')
            fonts.append(f"@font-face{{font-family:'{name}';font-weight:{weight};src:url({uri(file)})}}")
    parts=['<!doctype html><html><head><meta charset="utf-8"><title>Ecosphere — Molecular Attestation Research</title><style>'+''.join(fonts)+'''
    *{box-sizing:border-box}body{margin:0;background:#081d20}section{position:relative;width:1280px;height:720px;overflow:hidden;page-break-after:always;break-after:page;background:radial-gradient(125% 135% at 18% 12%,#15807a 0%,#0c5450 42%,#0a4043 70%,#0b2d31 100%)}
    .atom{position:absolute;margin:0} .txt{white-space:pre-wrap} @page{size:1280px 720px;margin:0} @media screen{section{margin:22px auto;box-shadow:0 5px 50px #0005}} @media print{section:last-child{break-after:auto}}
    </style></head><body>''']
    for i,s in enumerate(SL):
        parts.append(f'<section id="slide-{i+1}" aria-label="{html.escape(s["title"])}">')
        for a in s['items']:
            pos=f'left:{a["x"]}px;top:{a["y"]}px;width:{max(a["w"],1)}px;height:{max(a["h"],1)}px;'
            if a['kind']=='text':
                value=html.escape(a['value'])
                if a.get('href'):value=f'<a style="color:inherit;text-decoration:none" href="{html.escape(a["href"],quote=True)}">{value}</a>'
                parts.append(f'<div class="atom txt" style="{pos}font-family:\'{a["font"]}\';font-size:{a["size"]}px;font-weight:{700 if a["bold"] else 400};line-height:{a["line"]};color:#{a["color"]};text-align:{a["align"]};letter-spacing:{a["tracking"]}px">{value}</div>')
            elif a['kind']=='image':parts.append(f'<img class="atom" style="{pos}object-fit:{a["fit"]}" src="{uri(a["path"])}" alt="{Path(a["path"]).stem}">')
            elif a['kind'] in ('rect','circle'):
                border=f'border:1px solid #{a["stroke"]};' if a.get('stroke') else ''
                parts.append(f'<div class="atom" style="{pos}background:#{a["fill"]};{border}{"border-radius:50%;" if a["kind"]=="circle" else ""}"></div>')
            elif a['kind']=='line':
                style=f'border-top:{a["width"]}px solid #{a["color"]}' if a['w'] else f'border-left:{a["width"]}px solid #{a["color"]}'
                parts.append(f'<div class="atom" style="{pos}{style}"></div>')
        parts.append('</section>')
    parts.append('</body></html>');(OUT/'Ecosphere_Molecular_Attestation_2026.html').write_text(''.join(parts))


def pptx_deck():
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    scale=lambda p: Inches(p/96)
    for s in SL:
        sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.shapes.add_picture(str(AS/'gradient.png'),0,0,prs.slide_width,prs.slide_height)
        for a in s['items']:
            x,y,w,h=map(scale,(a['x'],a['y'],max(1,a['w']),max(1,a['h'])))
            if a['kind']=='image':
                from PIL import Image
                iw,ih=Image.open(a['path']).size
                if a['fit']=='contain':
                    ratio=min(a['w']/iw,a['h']/ih);nw,nh=iw*ratio,ih*ratio
                    sl.shapes.add_picture(a['path'],scale(a['x']+(a['w']-nw)/2),scale(a['y']+(a['h']-nh)/2),scale(nw),scale(nh))
                else:
                    sh=sl.shapes.add_picture(a['path'],x,y,w,h)
                    if iw/ih < a['w']/a['h']:
                        sh.crop_top=sh.crop_bottom=(1-(a['h']/a['w'])*(iw/ih))/2
                    else:sh.crop_left=sh.crop_right=(1-(a['w']/a['h'])/(iw/ih))/2
            elif a['kind']=='text':
                sh=sl.shapes.add_textbox(x,y,w,h);tf=sh.text_frame;tf.clear();tf.word_wrap=True
                tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
                for i,t in enumerate(a['value'].split('\n')):
                    p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=t
                    p.font.name=a['font'];p.font.size=Pt(a['size']*.75);p.font.bold=a['bold'];p.font.color.rgb=RGBColor.from_string(a['color'])
                    for run in p.runs:
                        if a['tracking']:run._r.get_or_add_rPr().set('spc',str(round(a['tracking']*75)))
                    p.line_spacing=a['line'];p.space_before=Pt(0);p.space_after=Pt(0)
                    p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[a['align']]
                # Shape-level links preserve the designed text color in Impress
                # as well as PowerPoint; text hyperlinks can force theme blue.
                if a.get('href'):sh.click_action.hyperlink.address=a['href']
            elif a['kind']=='line':
                sh=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,x,y,scale(a['x']+a['w']),scale(a['y']+a['h']));sh.line.color.rgb=RGBColor.from_string(a['color']);sh.line.width=Pt(a['width']*.75)
            else:
                sh=sl.shapes.add_shape(MSO_SHAPE.OVAL if a['kind']=='circle' else MSO_SHAPE.RECTANGLE,x,y,w,h)
                sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(a['fill']);sh.line.fill.background()
                for effect in sh._element.xpath('./p:style/a:effectRef'):effect.set('idx','0')
                if a.get('stroke'):sh.line.color.rgb=RGBColor.from_string(a['stroke']);sh.line.width=Pt(.75)
        sl.notes_slide.notes_text_frame.text=s['notes']
    prs.core_properties.title='Ecosphere — Molecular Attestation Intelligence'
    prs.core_properties.subject='MethaNet field-validation research proposal'
    prs.save(OUT/'Ecosphere_Molecular_Attestation_2026.pptx')


def render():
    (OUT/'slides').mkdir(exist_ok=True)
    with sync_playwright() as p:
        browser=p.chromium.launch(headless=True,args=['--no-sandbox']); page=browser.new_page(viewport={'width':1360,'height':820},device_scale_factor=2)
        page.goto((OUT/'Ecosphere_Molecular_Attestation_2026.html').as_uri());page.evaluate('document.fonts.ready')
        page.pdf(path=str(OUT/'Ecosphere_Molecular_Attestation_2026.pdf'),print_background=True,prefer_css_page_size=True)
        issues=page.evaluate('''() => [...document.querySelectorAll('section')].flatMap((s,i)=>[...s.querySelectorAll('.txt')].filter(e=>e.scrollHeight>e.clientHeight+2 || e.scrollWidth>e.clientWidth+2).map(e=>({slide:i+1,text:e.textContent,height:e.clientHeight,scroll:e.scrollHeight})))''')
        for i,el in enumerate(page.locator('section').all()):el.screenshot(path=str(OUT/f'slides/slide_{i+1:02}.png'))
        (OUT/'layout_audit.json').write_text(json.dumps({'slides':len(SL),'text_overflows':issues},indent=2))
        browser.close()
    from PIL import Image, ImageDraw
    canvas=Image.new('RGB',(1440,((len(SL)+2)//3)*290),'#091b20')
    for i in range(len(SL)):
        im=Image.open(OUT/f'slides/slide_{i+1:02}.png');im.thumbnail((480,270));canvas.paste(im,((i%3)*480,(i//3)*290))
        ImageDraw.Draw(canvas).text(((i%3)*480+9,(i//3)*290+272),f'{i+1:02}',fill='white')
    canvas.save(OUT/'contact_sheet.png')
    (OUT/'speaker_notes.md').write_text('# Ecosphere — speaker notes\n\n'+'\n\n'.join(f'## {i+1:02} · {s["title"].replace(chr(10)," ")}\n\n{s["notes"]}' for i,s in enumerate(SL)))
    print(json.dumps({'slides':len(SL),'overflows':issues,'output':str(OUT)},indent=2))
    if issues:raise SystemExit('Deck layout gate failed: text overflow')


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--draft',action='store_true',help='Permit explicit placeholders while computation is pending')
    args=parser.parse_args()
    OUT.mkdir(parents=True,exist_ok=True)
    if not args.draft:
        record=json.loads((OUT/'reconciliation_slide.json').read_text())
        provenance=json.loads((OUT/'figure_provenance.json').read_text())
        assert record['pilot_records']==662 and record['same_dna_controls']==23
        assert record['pilot_reproduction_gate']['pass']
        assert provenance['geometry_configuration']['status']=='verified'
        for entry in provenance['figures']:
            assert hashlib.sha256((OUT/entry['path']).read_bytes()).hexdigest()==entry['sha256']
        for name in ['atlas_reconciled.png','configuration_controls.png','evidence_graph_live.png']:
            assert (AS/name).is_file()
    # Transparent overlay is a layout element matching the original cover CSS.
    with sync_playwright() as p:
        b=p.chromium.launch(headless=True,args=['--no-sandbox']);pg=b.new_page(viewport={'width':1280,'height':720},device_scale_factor=2)
        pg.set_content('<body style="margin:0;width:1280px;height:720px;background:linear-gradient(95deg,rgba(9,40,44,.94) 0%,rgba(9,40,44,.82) 34%,rgba(9,40,44,.34) 66%,rgba(9,40,44,.12) 100%)"></body>')
        pg.screenshot(path=str(AS/'cover_overlay.png'),omit_background=True);b.close()
    make();html_deck();pptx_deck();render()
